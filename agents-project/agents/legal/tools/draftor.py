"""法律文书起草 — LLM 生成/修改 + docx 输出

支持两种场景：
1. 从零起草：根据用户需求生成新文书（内部 RAG 检索法律法规作为依据）
2. 基于已有文件修改/起草：从 RAG 按 file_id 读取上传文件内容（超长自动压缩）→ 重新生成
"""

import io
import json
import os
import re
import uuid
from datetime import datetime
from typing import Optional
from langchain.agents import create_agent
from langchain_core.messages import SystemMessage, HumanMessage, AIMessage
from langchain_core.tools import tool
from langgraph.config import get_config
from pydantic import BaseModel, Field
from agents.utils.models import MainModel, CompressModel, parse_json_object
from agents.utils.RAG import get_chunk, get_chunk_num
from agents.utils.executor import PREPROCESS_EXECUTOR, RAG_IO_EXECUTOR, run_in_pool
from .retrieve_law import retrieve_law_basis

try:
    from docx import Document
    from docx.shared import Pt, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
except ImportError:
    Document = None


DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


DRAFT_PROMPT = """你是一位资深法律文书起草专家。根据用户需求起草法律文书。只输出文书正文，不要输出任何解释、说明或建议。

## 文书类型
{doc_type}

## 用户需求
{requirement}

## 输出格式（严格遵守）
只输出一个 JSON 对象，禁止输出 markdown 代码块标记，禁止输出任何解释文字：
{{"title": "文书标题", "sections": [{{"heading": "条款标题，如 一、合同标的", "body": "该条款正文"}}]}}
其中 body 内如需多个自然段，用 \\n 分隔；全部条款按顺序放入 sections。

## 起草规则
1. 语言严谨、规范，符合法律文书格式
2. 条款编号清晰
3. 涉及金额、日期、期限等数字信息必须明确
4. 权利义务条款应当对等、公平
5. 必须包含管辖条款、送达条款
6. 起草前先调用 retrieve_law_basis 工具检索与本文书相关的法律条文（query 填文书主题与核心条款）；引用法律时必须只依据检索返回的条文，注明具体法律名称和条款；依据中未提供的，不得虚构具体法条编号，可用"根据相关法律规定"表述
7. 用户需求中未明确的部分，按最有利于保护用户利益的方向填充
8. 整篇文书应当完整、自洽
9. 禁止输出任何解释性文字，只输出文书正文本身
10. 文书内容必须使用中文"""

REVISE_PROMPT = """你是一位资深法律文书起草专家。根据修改意见对原文进行修订。只输出修订后的完整文书正文，不要输出任何解释、说明或建议。

## 文书类型
{doc_type}

## 原文
{original_content}

## 修改意见
{revision_instructions}

## 输出格式（严格遵守）
只输出一个 JSON 对象，禁止输出 markdown 代码块标记，禁止输出任何解释文字：
{{"title": "文书标题", "sections": [{{"heading": "条款标题，如 一、合同标的", "body": "该条款正文"}}]}}
其中 body 内如需多个自然段，用 \\n 分隔；输出完整修订后的文书，全部条款按顺序放入 sections。

## 修订规则
1. 逐条处理修改意见，确保每项都得到回应
2. 未涉及的部分保持原文不变
3. 修改后的条款需与全文其他条款协调一致
4. 涉及金额、日期、期限等数字信息必须精确
5. 输出完整文书，不得只输出修改的部分
6. 修订前先调用 retrieve_law_basis 工具检索与相关条款对应的法律条文；引用法律时必须只依据检索返回的条文，注明具体法律名称和条款；依据中未提供的，不得虚构具体法条编号，可用"根据相关法律规定"表述
7. 禁止输出任何解释性文字，只输出文书正文本身
8. 文书内容必须使用中文"""

# disable_streaming：图内 attach messages 流式回调时 invoke 会被 langchain 自动
# 改走流式，内部模型（需求压缩、起草 agent）的输出会作为 token 漏到前端。
_condense_model = CompressModel(temperature=0.1, disable_streaming=True)


class DraftSection(BaseModel):
    """文书中的一个条款/章节"""
    heading: str = Field(description='条款标题，如"一、合同标的"，不要带 markdown 标记')
    body: str = Field(description="该条款正文，可含多个自然段（用换行分隔）")


class DraftDocument(BaseModel):
    """结构化法律文书"""
    title: str = Field(description="文书标题")
    sections: list[DraftSection] = Field(description="文书全部条款，按顺序排列")


# checkpointer=False：默认 None 时子图会继承父图的 checkpointer 与 thread_id，把
# legal_agent 的整段对话 load 进来，连同"正在执行中、尚无 ToolMessage"的那条 tool_call
# 一起发给模型，触发 400 insufficient tool messages。子 agent 是单次无状态计算，必须切断继承。
# 不用 response_format（ToolStrategy/ProviderStrategy 都会 400）：DeepSeek V4 默认 thinking，
# thinking 模式直接拒绝 tool_choice 与 response_format=json_schema。改为提示词约束模型只输出
# JSON，再由 _parse_structured 解析。工具循环仍保留，故 recursion_limit 需抬高。
_draft_agent = create_agent(
    model=MainModel(temperature=0.3, disable_streaming=True),
    tools=[retrieve_law_basis],
    checkpointer=False,
)


def _last_agent_text(result: dict) -> str:
    """从 create_agent 结果里取最终无工具调用的 AI 文本。"""
    messages = result.get("messages", [])
    for msg in reversed(messages):
        if isinstance(msg, AIMessage) and not getattr(msg, "tool_calls", None):
            content = msg.content
            if content:
                return content if isinstance(content, str) else str(content)
    return "" if not messages else str(messages[-1].content)

# 原始文件内容超过该长度时，先 LLM 压缩成要点摘要再注入提示词，避免超长上下文
FULL_TEXT_LIMIT = 6000
CONDENSE_INPUT_LIMIT = 30000


CONDENSE_PROMPT = """你是一位法律文书分析助手。请将下面这份文书的原文压缩成结构化的要点摘要，
供后续起草/修改文书时作为参考。只输出摘要本身，不要任何解释。

## 保留内容
- 当事人/合同主体、合同标的与内容
- 关键金额、日期、期限等数字信息
- 权利义务要点、违约责任、争议解决、管辖、送达等关键条款（保留条款编号）
- 容易被忽略但对修改重要的细节条款

## 要求
1. 用简洁条目式输出，保留原文条款编号
2. 忽略格式性套话、重复表述
3. 总输出不超过 1500 字
4. 只输出摘要，禁止解释性文字

## 文书原文
{content}"""


def _clean_filename(name: str) -> str:
    return re.sub(r'[\\/:*?"<>|]', "_", name.strip())[:50]


def _load_file_text(file_id: str) -> str:
    """从 RAG 按 file_id 取回上传文件全部切片，并按顺序拼接成文本。"""
    total = get_chunk_num(file_id)
    if total == 0:
        return ""
    chunks = get_chunk(file_id, page=1, limit=total)
    return "\n".join(c["content"] for c in chunks)


async def _prepare_original_content(file_id: str) -> str:
    """取回上传文件文本；超长时先 LLM 压缩成要点摘要，控制注入提示词的长度。"""
    text = await run_in_pool(RAG_IO_EXECUTOR, _load_file_text, file_id)
    if not text:
        return ""
    if len(text) <= FULL_TEXT_LIMIT:
        return text
    try:
        resp = await _condense_model.ainvoke([
            SystemMessage(
                content=CONDENSE_PROMPT.format(content=text[:CONDENSE_INPUT_LIMIT])
            ),
        ])
        condensed = (resp.content or "").strip()
        return condensed if condensed else text[:FULL_TEXT_LIMIT]
    except Exception:
        return text[:FULL_TEXT_LIMIT]


def _build_docx(title: str, sections: list[tuple[str, str]]) -> Document:
    """将结构化内容生成为 .docx。"""
    doc = Document()

    for section in doc.sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(3.17)
        section.right_margin = Cm(3.17)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p_title.add_run(title)
    run.bold = True
    run.font.size = Pt(18)
    run.font.name = "黑体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")

    for heading, body in sections:
        p_head = doc.add_paragraph()
        run = p_head.add_run(heading)
        run.bold = True
        run.font.size = Pt(14)
        run.font.name = "黑体"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")

        for para_text in body.split("\n"):
            para_text = para_text.strip()
            if not para_text:
                continue
            p = doc.add_paragraph()
            run = p.add_run(para_text)
            run.font.size = Pt(12)
            run.font.name = "仿宋"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "仿宋")
            p.paragraph_format.first_line_indent = Pt(24)
            p.paragraph_format.line_spacing = Pt(28)

    return doc


def _parse_llm_output(text: str) -> tuple[str, list[tuple[str, str]]]:
    """解析 LLM 输出为标题和段落列表（结构化输出失败时的兜底）。"""
    title = "法律文书"
    sections: list[tuple[str, str]] = []
    lines = text.strip().split("\n")
    current_heading = None
    current_body: list[str] = []

    def _flush():
        if current_heading:
            sections.append((current_heading, "\n".join(current_body).strip()))

    for line in lines:
        line = line.strip()
        if not line:
            continue
        # 先剥掉行首 markdown 标题前缀，避免 "## TITLE: x" 被当成章节而丢掉标题
        norm = line[3:].strip() if line.startswith("## ") else line
        upper = norm.upper()
        if upper.startswith("TITLE:"):
            title = norm[len("TITLE:"):].strip()
        elif upper.startswith("SECTION:"):
            _flush()
            current_heading = norm[len("SECTION:"):].strip()
            current_body = []
        else:
            current_body.append(line)

    _flush()
    return title, sections


def _parse_structured(text: str) -> tuple[str, list[tuple[str, str]]]:
    """把模型输出解析为 (标题, [(条款标题, 正文)])。

    首选 JSON（提示词要求只输出 JSON），失败再退回 TITLE/SECTION 文本解析。
    """
    try:
        doc = DraftDocument.model_validate(parse_json_object(text))
        if doc.sections:
            return doc.title, [(s.heading, s.body) for s in doc.sections]
        print("[draftor] JSON 解析成功但 sections 为空，回退文本解析")
    except Exception as e:
        print(f"[draftor] JSON 解析失败，回退文本解析: {e}")
    return _parse_llm_output(text)


def _render_docx(title: str, sections: list[tuple[str, str]]) -> tuple[str, bytes]:
    """同步：生成 .docx 字节和文件名（python-docx 是 CPU 活，放专用池调用）。"""
    doc = _build_docx(title, sections)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"{_clean_filename(title)}_{timestamp}.docx"
    buf = io.BytesIO()
    doc.save(buf)
    return filename, buf.getvalue()


async def draft_document(
    doc_type: str,
    requirement: str,
    revision_instructions: Optional[str] = None,
    original_file_id: Optional[str] = None,
) -> tuple[str, bytes]:
    """起草或修改法律文书，生成 .docx 字节。

    Args:
        doc_type: 文书类型，如：合同、律师函、起诉状、协议书、声明书
        requirement: 用户需求描述（修改场景下为修改后的目标需求）
        revision_instructions: 修改意见（基于上传文件修改时传入）
        original_file_id: 用户上传文件在 RAG 中的 file_id（基于上传文件修改/起草时传入）

    Returns:
        (文件名, .docx 文件字节)
    """
    if Document is None:
        raise ImportError("python-docx 未安装，无法生成 .docx 文件")

    # 区分"初次生成"与"基于上传文件修改/起草"：
    # 用户上传文件不落盘，只以 file_id 存在 RAG 中，因此只走 original_file_id 这条路。
    original_content = None
    if original_file_id:
        original_content = await _prepare_original_content(original_file_id)
        if not original_content:
            raise ValueError("未能在 RAG 中找到该文件内容，文件可能已过期，请重新上传后再试。")

    if original_content:
        # 基于已有文件：revision_instructions 为空时，把 requirement 当作起草/修改需求
        rev_instructions = revision_instructions or requirement
        system_prompt = REVISE_PROMPT.format(
            doc_type=doc_type,
            original_content=original_content,
            revision_instructions=rev_instructions,
        )
        human_content = f"请基于这份{doc_type}按要求起草/修订：{rev_instructions}"
    else:
        # 初次生成
        system_prompt = DRAFT_PROMPT.format(
            doc_type=doc_type,
            requirement=requirement,
        )
        human_content = f"请起草一份{doc_type}，需求：{requirement}"

    # 法律依据由 agent 起草/修订前通过 retrieve_law 工具自行检索；最终输出为 JSON 文本。
    result = await _draft_agent.ainvoke(
        {
            "messages": [
                SystemMessage(content=system_prompt),
                HumanMessage(content=human_content),
            ]
        },
        config={"recursion_limit": 20},
    )

    title, sections = _parse_structured(_last_agent_text(result))

    return await run_in_pool(PREPROCESS_EXECUTOR, _render_docx, title, sections)


def _save_created_file(filename: str, data: bytes) -> str:
    """把生成的文书字节写入 PG legal_created_file，返回行 id。

    userId/threadId 取自 graph 运行 config，与 temp_file / 会话表的隔离 key 一致。
    """
    import psycopg

    config = get_config().get("configurable", {})
    user_id = config.get("userId")
    thread_id = config.get("thread_id") or ""
    if user_id is None:
        raise RuntimeError("缺少 userId，无法保存生成的文书")

    conninfo = os.getenv("postgreSQL_URL")
    if not conninfo:
        raise RuntimeError("缺少 postgreSQL_URL，无法保存生成的文书")

    file_id = uuid.uuid4().hex
    with psycopg.connect(conninfo) as conn:
        conn.execute(
            """
            INSERT INTO legal_created_file
                (id, "userId", "threadId", "originalName", "mimeType", "sizeBytes", data)
            VALUES (%s, %s, %s, %s, %s, %s, %s)
            """,
            (file_id, str(user_id), thread_id, filename, DOCX_MIME, len(data), data),
        )
        conn.commit()
    return file_id


@tool
async def draft_document_agent(
    doc_type: str,
    requirement: str,
    revision_instructions: Optional[str] = None,
    original_file_id: Optional[str] = None,
) -> str:
    """起草或修改法律文书，生成 .docx 并入库，返回 JSON（含 file_id 与 filename）。

    支持两种使用方式：
    1. 初次生成：传入 doc_type + requirement，不传 original_file_id
    2. 基于上传文件修改/起草：传入 doc_type + requirement + original_file_id（用户上传文件在
       RAG 中的 file_id，即调度状态里的"已解析文件ID"），具体修改/起草需求填 revision_instructions

    Args:
        doc_type: 文书类型，如：合同、律师函、起诉状、协议书、声明书
        requirement: 用户需求描述
        revision_instructions: 修改/起草意见
        original_file_id: 用户上传文件在 RAG 中的 file_id（基于上传文件修改/起草时必填）
    """
    print("文件生成工具被调用了")
    filename, data = await draft_document(doc_type, requirement, revision_instructions, original_file_id)
    file_id = await run_in_pool(RAG_IO_EXECUTOR, _save_created_file, filename, data)
    return json.dumps({"file_id": file_id, "filename": filename}, ensure_ascii=False)
