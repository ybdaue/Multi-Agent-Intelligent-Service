from __future__ import annotations

import re
from typing import TYPE_CHECKING, Any, Dict, List

if TYPE_CHECKING:
    from .multiAgent import FileSource


TEXT_SUFFIXES = {".txt", ".md", ".json", ".csv", ".xml", ".html", ".htm"}
PDF_SUFFIXES = {".pdf"}
WORD_SUFFIXES = {".docx", ".doc"}
EXCEL_SUFFIXES = {".xlsx", ".xls"}

DEFAULT_CHUNK_SIZE = 700
DEFAULT_OVERLAP = 150



def _parse_txt(raw: bytes) -> str:
    return raw.decode("utf-8", errors="replace")

def _parse_pdf(raw: bytes) -> str:
    import pdfplumber
    import io

    parts: list[str] = []
    with pdfplumber.open(io.BytesIO(raw)) as pdf:
        for page in pdf.pages:
            text = page.extract_text()
            if text:
                parts.append(text)
    return "\n".join(parts)

def _parse_docx(raw: bytes) -> str:
    import docx
    import io

    doc = docx.Document(io.BytesIO(raw))
    parts: list[str] = []

    for p in doc.paragraphs:
        if p.text.strip():
            parts.append(p.text)

    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                parts.append(row_text)

    return "\n".join(parts)


def _parse_doc(raw: bytes, file_source: FileSource) -> str:
    """Parse a .doc file (legacy OLE format).

    Tries docx parser first (some files are mislabeled), then antiword.
    """
    # Try docx parser first — some .doc files are actually .docx format
    try:
        return _parse_docx(raw)
    except Exception:
        pass

    # If path is available, try antiword
    if file_source["type"] == "path":
        import shutil
        import subprocess

        antiword = shutil.which("antiword")
        if antiword:
            try:
                result = subprocess.run(
                    [antiword, "-w", "0", str(file_source["content"])],
                    capture_output=True,
                    text=True,
                    timeout=30,
                )
                if result.returncode == 0 and result.stdout.strip():
                    return result.stdout.strip()
            except Exception:
                pass

    raise ValueError(
        "无法解析 .doc 文件。请使用 Microsoft Word 将文件另存为 .docx 格式后重新上传。"
    )

def _parse_xlsx(raw: bytes) -> str:
    import openpyxl
    import io

    wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True)
    parts: list[str] = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows: list[str] = []
        for row in ws.iter_rows(values_only=True):
            line = " | ".join(str(c) if c is not None else "" for c in row).strip()
            if line:
                rows.append(line)
        if rows:
            parts.append(f"[{sheet_name}]")
            parts.extend(rows)
    return "\n".join(parts)

def _parse_xls(raw: bytes) -> str:
    import xlrd
    import io

    wb = xlrd.open_workbook(file_contents=raw)
    parts: list[str] = []
    for sheet_name in wb.sheet_names():
        ws = wb.sheet_by_name(sheet_name)
        rows: list[str] = []
        for row_idx in range(ws.nrows):
            cells = [str(ws.cell_value(row_idx, c)) for c in range(ws.ncols)]
            line = " | ".join(c.strip() for c in cells if c.strip()).strip()
            if line:
                rows.append(line)
        if rows:
            parts.append(f"[{sheet_name}]")
            parts.extend(rows)
    return "\n".join(parts)


#文件解析方法
def parse_file(file_source: FileSource) -> str:
    """Parse an uploaded file into plain text.

    Supports: .txt / .pdf / .docx / .doc / .xlsx / .xls
    """
    suffix = file_source["suffix"].lower()

    if file_source["type"] == "path":
        path = file_source["content"]
        if not isinstance(path, str):
            path = str(path)
        with open(path, "rb") as f:
            raw = f.read()
    else:
        raw = file_source["content"]
        if isinstance(raw, str):
            raw = raw.encode("utf-8")

    if suffix in TEXT_SUFFIXES:
        return _parse_txt(raw)
    if suffix in PDF_SUFFIXES:
        return _parse_pdf(raw)
    if suffix == ".docx":
        return _parse_docx(raw)
    if suffix == ".doc":
        return _parse_doc(raw, file_source)
    if suffix == ".xls":
        return _parse_xls(raw)
    if suffix in EXCEL_SUFFIXES:
        return _parse_xlsx(raw)

    raise ValueError(f"暂不支持的文件类型: {suffix}")

_RE_CHAPTER = re.compile(r"第[一二三四五六七八九十百千\d]+章")
_RE_ARTICLE = re.compile(r"第[一二三四五六七八九十百千\d]+条")

_RE_EN_ARTICLE = re.compile(r"Article\s+\d+", re.IGNORECASE)


#根据关键字出现次数，判断这是什么类型的文件
def _detect_structure(text: str) -> str:
    """Detect document structure type for chunking decisions.

    Returns ``"article"``, ``"chapter"``, or ``"paragraph"``.
    """
    article_count = len(_RE_ARTICLE.findall(text))
    chapter_count = len(_RE_CHAPTER.findall(text))
    en_article_count = len(_RE_EN_ARTICLE.findall(text))

    if article_count >= 3 or en_article_count >= 3:
        return "article"
    if chapter_count >= 2:
        return "chapter"
    return "paragraph"

#决定切分策略，文章按条或章切分
def _split_structured(text: str) -> List[Dict[str, Any]]:
    """Split text by **条** or **Article** markers, preserving headings."""
    markers = _RE_ARTICLE.findall(text)
    parts = _RE_ARTICLE.split(text)

    if len(markers) < 3:
        markers = _RE_EN_ARTICLE.findall(text)
        parts = _RE_EN_ARTICLE.split(text)

    if len(markers) < 3:
        markers = _RE_CHAPTER.findall(text)
        parts = _RE_CHAPTER.split(text)

    result: List[Dict[str, Any]] = []

    if not markers:
        return result

    preamble = parts[0].strip()
    if preamble:
        result.append({"heading": None, "content": preamble})

    for i, marker in enumerate(markers):
        body = (parts[i + 1] if i + 1 < len(parts) else "").strip()
        result.append({
            "heading": marker,
            "content": f"{marker} {body}" if body else marker,
        })

    return result

#把太短的相邻的chunk合并起来
def _merge_small_chunks(
    chunks: List[Dict[str, Any]],
    *,
    min_size: int = 200,
) -> List[Dict[str, Any]]:
    """Merge adjacent small chunks so individual pieces don't lose context."""
    if not chunks:
        return []

    merged: List[Dict[str, Any]] = []
    buf = dict(chunks[0])

    for cur in chunks[1:]:
        buf_len = len(buf["content"])

        if buf_len < min_size:
            # Merge current into buffer
            buf["content"] += "\n\n" + cur["content"]
        elif len(cur["content"]) < min_size and cur.get("heading") and buf.get("heading"):
            # Current is a small headed unit (e.g. a short article)
            buf["content"] += "\n\n" + cur["content"]
        else:
            merged.append(buf)
            buf = dict(cur)

    merged.append(buf)
    return merged

#如果一个 chunk 超过了max_size，就进一步细分
def _split_oversized(
    chunks: List[Dict[str, Any]],
    *,
    max_size: int,
    overlap: int,
) -> List[Dict[str, Any]]:
    """Further split any chunk exceeding ``max_size``."""
    result: List[Dict[str, Any]] = []

    for chunk in chunks:
        text = chunk["content"]
        if len(text) <= max_size:
            result.append(chunk)
            continue

        # Level 1: split by double-newline paragraphs
        paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]

        if len(paragraphs) > 1:
            segments = _overlap_segments(paragraphs, max_size, overlap, sep="\n\n")
        else:
            # Level 2: split by sentence
            sentences = re.split(r"(?<=[。！？；.!?;])\s*", text)
            sentences = [s.strip() for s in sentences if s.strip()]
            segments = _overlap_segments(sentences, max_size, overlap, sep="")

        for i, seg in enumerate(segments):
            heading = chunk["heading"]
            if i > 0 and heading:
                heading = f"{heading}(续)"
            result.append({"heading": heading, "content": seg})

    return result

#滑动窗口，打包后回退overlap个字符
def _overlap_segments(
    units: List[str],
    max_size: int,
    overlap: int,
    sep: str,
) -> List[str]:
    """Build overlapping segments from a list of text units."""
    segments: List[str] = []
    i = 0

    while i < len(units):
        seg = units[i]
        j = i + 1
        while j < len(units) and len(seg) + len(sep) + len(units[j]) <= max_size:
            seg += sep + units[j]
            j += 1

        segments.append(seg)

        overlap_chars = 0
        k = j - 1
        while k >= i and overlap_chars < overlap:
            overlap_chars += len(units[k]) + len(sep)
            k -= 1
        i = max(i + 1, k + 1)

    return segments


#切片方法
def chunk_text(
    text: str,
    file_id: str,
    *,
    chunk_size: int = DEFAULT_CHUNK_SIZE,
    overlap: int = DEFAULT_OVERLAP,
) -> List[Dict[str, Any]]:
    """Hierarchical chunking for legal documents (律师函 / 合同 / 协议).

    Strategy
    --------
    1. **Structure detection** – determine whether the document is
       article-based (第X条 / Article N), chapter-based (第X章), or
       free-form prose.
    2. **Primary split** – break at the most granular structural boundary
       so each chunk is a semantically complete unit.
    3. **Merge** – group adjacent tiny units that would lose meaning alone.
    4. **Oversized split** – further divide any chunk exceeding
       ``chunk_size``, first by paragraph then by sentence, with
       ``overlap`` characters of context carried over.
    5. **Metadata** – every chunk carries ``file_id``, ``chunk_index``,
       ``total_chunks``, and the structural ``heading`` for downstream
       RAG retrieval.

    Returns (list of dict)
        [
            {
                "content": "第十二条 违约责任 ...",
                "metadata": {
                    "file_id": "xxx",
                    "chunk_index": 3,
                    "total_chunks": 12,
                    "heading": "第十二条",
                }
            },
            ...
        ]
    """
    text = text.strip()
    if not text:
        return []

    structure = _detect_structure(text)

    # ── 1. Primary split ──
    if structure in ("article", "chapter"):
        raw = _split_structured(text)
    else:
        # Paragraph-based
        paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]
        raw = [{"heading": None, "content": p} for p in paragraphs]

    if not raw:
        raw = [{"heading": None, "content": text}]

    # ── 2. Merge small chunks ──
    merged = _merge_small_chunks(raw, min_size=chunk_size // 3)

    # ── 3. Split oversized ──
    final = _split_oversized(merged, max_size=chunk_size, overlap=overlap)

    # ── 4. Attach metadata ──
    total = len(final)
    return [
        {
            "content": chunk["content"],
            "metadata": {
                "file_id": file_id,
                "chunk_index": idx,
                "total_chunks": total,
                "heading": chunk["heading"],
            },
        }
        for idx, chunk in enumerate(final)
    ]

#一站式，解析+切片
def parse_and_chunk(
    file_source: FileSource,
    *,
    chunk_size: int = DEFAULT_CHUNK_SIZE,
    overlap: int = DEFAULT_OVERLAP,
) -> List[Dict[str, Any]]:
    """Convenience: parse a file, then chunk the result."""
    text = parse_file(file_source)
    return chunk_text(text, file_source["file_id"], chunk_size=chunk_size, overlap=overlap)
