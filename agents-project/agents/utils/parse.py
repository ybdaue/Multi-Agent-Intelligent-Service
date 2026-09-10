from __future__ import annotations

import re
from typing import TYPE_CHECKING, Any, Dict, List

if TYPE_CHECKING:
    from agents.multiAgent import FileSource


TEXT_SUFFIXES = {".txt", ".md", ".json", ".csv", ".xml", ".html", ".htm"}
PDF_SUFFIXES = {".pdf"}
WORD_SUFFIXES = {".docx", ".doc"}
EXCEL_SUFFIXES = {".xlsx", ".xls"}

DEFAULT_CHUNK_SIZE = 700
DEFAULT_OVERLAP = 150



def _parse_txt(raw: bytes) -> str:
    return raw.decode("utf-8", errors="replace")

def _parse_pdf(raw: bytes) -> str:
    import pdfplumber
    import io

    parts: list[str] = []
    with pdfplumber.open(io.BytesIO(raw)) as pdf:
        for page in pdf.pages:
            text = page.extract_text()
            if text:
                parts.append(text)
    return "\n".join(parts)

def _parse_docx(raw: bytes) -> str:
    import docx
    import io

    doc = docx.Document(io.BytesIO(raw))
    parts: list[str] = []

    for p in doc.paragraphs:
        if p.text.strip():
            parts.append(p.text)

    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                parts.append(row_text)

    return "\n".join(parts)


def _parse_doc(raw: bytes, file_source: FileSource) -> str:
    """Parse a .doc file (legacy OLE format).

    Tries docx parser first (some files are mislabeled), then antiword.
    """
    # Try docx parser first — some .doc files are actually .docx format
    try:
        return _parse_docx(raw)
    except Exception:
        pass

    # If path is available, try antiword
    if file_source["type"] == "path":
        import shutil
        import subprocess

        antiword = shutil.which("antiword")
        if antiword:
            try:
                result = subprocess.run(
                    [antiword, "-w", "0", str(file_source["content"])],
                    capture_output=True,
                    text=True,
                    timeout=30,
                )
                if result.returncode == 0 and result.stdout.strip():
                    return result.stdout.strip()
            except Exception:
                pass

    raise ValueError(
        "无法解析 .doc 文件。请使用 Microsoft Word 将文件另存为 .docx 格式后重新上传。"
    )

def _parse_xlsx(raw: bytes) -> str:
    import openpyxl
    import io

    wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True)
    parts: list[str] = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows: list[str] = []
        for row in ws.iter_rows(values_only=True):
            line = " | ".join(str(c) if c is not None else "" for c in row).strip()
            if line:
                rows.append(line)
        if rows:
            parts.append(f"[{sheet_name}]")
            parts.extend(rows)
    return "\n".join(parts)

def _parse_xls(raw: bytes) -> str:
    import xlrd
    import io

    wb = xlrd.open_workbook(file_contents=raw)
    parts: list[str] = []
    for sheet_name in wb.sheet_names():
        ws = wb.sheet_by_name(sheet_name)
        rows: list[str] = []
        for row_idx in range(ws.nrows):
            cells = [str(ws.cell_value(row_idx, c)) for c in range(ws.ncols)]
            line = " | ".join(c.strip() for c in cells if c.strip()).strip()
            if line:
                rows.append(line)
        if rows:
            parts.append(f"[{sheet_name}]")
            parts.extend(rows)
    return "\n".join(parts)


def _load_uploaded_bytes(temp_file_id: str) -> bytes:
    """按 temp_file 行 id 同步读回大文件字节（切片在 preprocessing 同步节点内进行）。"""
    import os

    import psycopg

    conninfo = os.getenv("postgreSQL_URL")
    if not conninfo:
        raise RuntimeError("缺少 postgreSQL_URL，无法按 id 读取上传文件")
    with psycopg.connect(conninfo) as conn:
        with conn.cursor() as cur:
            cur.execute('SELECT data FROM temp_file WHERE id = %s', (temp_file_id,))
            row = cur.fetchone()
    if row is None:
        raise ValueError("上传文件不存在或已过期，请重新上传")
    return bytes(row[0])


#文件解析方法
def parse_file(file_source: FileSource) -> str:
    """Parse an uploaded file into plain text.

    Supports: .txt / .pdf / .docx / .doc / .xlsx / .xls
    content 字段：path=磁盘路径 / db=temp_file 行 id / 其余按 bytes 处理
    """
    suffix = file_source["suffix"].lower()

    if file_source["type"] == "path":
        path = file_source["content"]
        if not isinstance(path, str):
            path = str(path)
        with open(path, "rb") as f:
            raw = f.read()
    elif file_source["type"] == "db":
        raw = _load_uploaded_bytes(file_source["content"])
    else:
        raw = file_source["content"]
        if isinstance(raw, str):
            raw = raw.encode("utf-8")

    if suffix in TEXT_SUFFIXES:
        return _parse_txt(raw)
    if suffix in PDF_SUFFIXES:
        return _parse_pdf(raw)
    if suffix == ".docx":
        return _parse_docx(raw)
    if suffix == ".doc":
        return _parse_doc(raw, file_source)
    if suffix == ".xls":
        return _parse_xls(raw)
    if suffix in EXCEL_SUFFIXES:
        return _parse_xlsx(raw)

    raise ValueError(f"暂不支持的文件类型: {suffix}")

_RE_CHAPTER = re.compile(r"第[一二三四五六七八九十百千\d]+章")
_RE_ARTICLE = re.compile(r"第[一二三四五六七八九十百千\d]+条")

_RE_EN_ARTICLE = re.compile(r"Article\s+\d+", re.IGNORECASE)


#根据关键字出现次数，判断这是什么类型的文件
def _detect_structure(text: str) -> str:
    """Detect document structure type for chunking decisions.

    Returns ``"article"``, ``"chapter"``, or ``"paragraph"``.
    """
    article_count = len(_RE_ARTICLE.findall(text))
    chapter_count = len(_RE_CHAPTER.findall(text))
    en_article_count = len(_RE_EN_ARTICLE.findall(text))

    if article_count >= 3 or en_article_count >= 3:
        return "article"
    if chapter_count >= 2:
        return "chapter"
    return "paragraph"

#决定切分策略，文章按条或章切分
def _split_structured(text: str) -> List[Dict[str, Any]]:
    """Split text by **条** or **Article** markers, preserving headings."""
    markers = _RE_ARTICLE.findall(text)
    parts = _RE_ARTICLE.split(text)

    if len(markers) < 3:
        markers = _RE_EN_ARTICLE.findall(text)
        parts = _RE_EN_ARTICLE.split(text)

    if len(markers) < 3:
        markers = _RE_CHAPTER.findall(text)
        parts = _RE_CHAPTER.split(text)

    result: List[Dict[str, Any]] = []

    if not markers:
        return result

    preamble = parts[0].strip()
    if preamble:
        result.append({"heading": None, "content": preamble})

    for i, marker in enumerate(markers):
        body = (parts[i + 1] if i + 1 < len(parts) else "").strip()
        result.append({
            "heading": marker,
            "content": f"{marker} {body}" if body else marker,
        })

    return result

#把太短的相邻的chunk合并起来
def _merge_small_chunks(
    chunks: List[Dict[str, Any]],
    *,
    min_size: int = 200,
) -> List[Dict[str, Any]]:
    """Merge adjacent small chunks so individual pieces don't lose context."""
    if not chunks:
        return []

    merged: List[Dict[str, Any]] = []
    buf = dict(chunks[0])

    for cur in chunks[1:]:
        buf_len = len(buf["content"])

        if buf_len < min_size:
            # Merge current into buffer
            buf["content"] += "\n\n" + cur["content"]
        elif len(cur["content"]) < min_size and cur.get("heading") and buf.get("heading"):
            # Current is a small headed unit (e.g. a short article)
            buf["content"] += "\n\n" + cur["content"]
        else:
            merged.append(buf)
            buf = dict(cur)

    merged.append(buf)
    return merged

#如果一个 chunk 超过了max_size，就进一步细分
def _split_oversized(
    chunks: List[Dict[str, Any]],
    *,
    max_size: int,
    overlap: int,
) -> List[Dict[str, Any]]:
    """Further split any chunk exceeding ``max_size``."""
    result: List[Dict[str, Any]] = []

    for chunk in chunks:
        text = chunk["content"]
        if len(text) <= max_size:
            result.append(chunk)
            continue

        # Level 1: split by double-newline paragraphs
        paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]

        if len(paragraphs) > 1:
            segments = _overlap_segments(paragraphs, max_size, overlap, sep="\n\n")
        else:
            # Level 2: split by sentence
            sentences = re.split(r"(?<=[。！？；.!?;])\s*", text)
            sentences = [s.strip() for s in sentences if s.strip()]
            segments = _overlap_segments(sentences, max_size, overlap, sep="")

        for i, seg in enumerate(segments):
            heading = chunk["heading"]
            if i > 0 and heading:
                heading = f"{heading}(续)"
            result.append({"heading": heading, "content": seg})

    return result

#滑动窗口，打包后回退overlap个字符
def _overlap_segments(
    units: List[str],
    max_size: int,
    overlap: int,
    sep: str,
) -> List[str]:
    """Build overlapping segments from a list of text units."""
    segments: List[str] = []
    i = 0

    while i < len(units):
        seg = units[i]
        j = i + 1
        while j < len(units) and len(seg) + len(sep) + len(units[j]) <= max_size:
            seg += sep + units[j]
            j += 1

        segments.append(seg)

        overlap_chars = 0
        k = j - 1
        while k >= i and overlap_chars < overlap:
            overlap_chars += len(units[k]) + len(sep)
            k -= 1
        i = max(i + 1, k + 1)

    return segments


#切片方法
def chunk_text(
    text: str,
    file_id: str,
    *,
    chunk_size: int = DEFAULT_CHUNK_SIZE,
    overlap: int = DEFAULT_OVERLAP,
) -> List[Dict[str, Any]]:
    """Hierarchical chunking for legal documents (律师函 / 合同 / 协议).

    Strategy
    --------
    1. **Structure detection** – determine whether the document is
       article-based (第X条 / Article N), chapter-based (第X章), or
       free-form prose.
    2. **Primary split** – break at the most granular structural boundary
       so each chunk is a semantically complete unit.
    3. **Merge** – group adjacent tiny units that would lose meaning alone.
    4. **Oversized split** – further divide any chunk exceeding
       ``chunk_size``, first by paragraph then by sentence, with
       ``overlap`` characters of context carried over.
    5. **Metadata** – every chunk carries ``file_id``, ``chunk_index``,
       ``total_chunks``, and the structural ``heading`` for downstream
       RAG retrieval.

    Returns (list of dict)
        [
            {
                "content": "第十二条 违约责任 ...",
                "metadata": {
                    "file_id": "xxx",
                    "chunk_index": 3,
                    "total_chunks": 12,
                    "heading": "第十二条",
                }
            },
            ...
        ]
    """
    text = text.strip()
    if not text:
        return []

    structure = _detect_structure(text)

    # ── 1. Primary split ──
    if structure in ("article", "chapter"):
        raw = _split_structured(text)
    else:
        # Paragraph-based
        paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]
        raw = [{"heading": None, "content": p} for p in paragraphs]

    if not raw:
        raw = [{"heading": None, "content": text}]

    # ── 2. Merge small chunks ──
    merged = _merge_small_chunks(raw, min_size=chunk_size // 3)

    # ── 3. Split oversized ──
    final = _split_oversized(merged, max_size=chunk_size, overlap=overlap)

    # ── 4. Attach metadata ──
    total = len(final)
    return [
        {
            "content": chunk["content"],
            "metadata": {
                "file_id": file_id,
                "chunk_index": idx,
                "total_chunks": total,
                "heading": chunk["heading"],
            },
        }
        for idx, chunk in enumerate(final)
    ]

# =============================================================================
# 法律库父子文档切分（仅 legal_ingest 使用；上传文档仍走 chunk_text）
# 父模块 = 结构章/条群组，约 parent_min..parent_max 字符；子模块 = 父内约 child_max 的检索单元。
# 子存 Chroma + bm25（metadata/行带父信息），父存 Postgres laws 表。检索用子召回、返回父全文。
# =============================================================================

LAW_PARENT_MIN = 800
LAW_PARENT_MAX = 2000
LAW_CHILD_MAX = 300

_RE_CHAPTER_FULL = re.compile(r"第[一二三四五六七八九十百千\d]+章")
_RE_ARTICLE_ONLY = re.compile(r"(第[一二三四五六七八九十百千\d]+条)")

# 取正文首行作为法规名（heading 用）；若首行是结构标记则退回 file_id
def _law_title(text: str, file_id: str) -> str:
    first = (text.split("\n", 1)[0] or "").strip()
    if first and not _RE_CHAPTER_FULL.match(first) and not _RE_ARTICLE_ONLY.match(first):
        return first[:60]
    return file_id

# 剥离法规名/日期/目录等前部内容，返回从正文第一章起的文本
def _law_body(text: str) -> str:
    chapters = [m.start() for m in _RE_CHAPTER_FULL.finditer(text)]
    articles = [m.start() for m in _RE_ARTICLE_ONLY.finditer(text)]

    if chapters:
        # 真实正文章：该章之后、下一章之前存在条款
        body_start = None
        for ci in range(len(chapters)):
            seg_start = chapters[ci]
            seg_end = chapters[ci + 1] if ci + 1 < len(chapters) else len(text)
            if any(seg_start < a < seg_end for a in articles):
                body_start = seg_start
                break
        if body_start is None:
            body_start = chapters[0]
        return text[body_start:].strip()

    if articles:
        return text[articles[0]:].strip()

    return text.strip()

# 按真实章行切分成若干章块；无章则整篇一块。块首保留章标题行。
def _split_macro_blocks(body: str) -> list[str]:
    if not body:
        return []
    if not _RE_CHAPTER_FULL.search(body):
        return [body]
    starts = [m.start() for m in _RE_CHAPTER_FULL.finditer(body)]
    blocks = []
    for i, s in enumerate(starts):
        e = starts[i + 1] if i + 1 < len(starts) else len(body)
        blocks.append(body[s:e].strip())
    return [b for b in blocks if b]

# 把整条单元顺序打包成 ≤ parent_max 的父文本
def _pack_to_parents(units: list[dict], parent_max: int) -> list[str]:
    parents: list[str] = []
    buf = ""
    for u in units:
        u_text = (u.get("text") or "").strip()
        if not u_text:
            continue
        if buf:
            if len(buf) + len(u_text) + 2 <= parent_max:
                buf = f"{buf}\n\n{u_text}"
                continue
            parents.append(buf)
        buf = u_text
    if buf:
        parents.append(buf)
    return parents

# 产出父模块列表 [{content, heading, parent_index}]
def _build_law_parents(body: str, law_title: str, parent_min: int, parent_max: int) -> list[dict]:
    parents: list[dict] = []
    for block in _split_macro_blocks(body):
        units = _split_article_units(block)
        if not units:
            units = [{"heading": None, "text": block}]
        for content in _pack_to_parents(units, parent_max):
            parents.append({
                "content": content,
                "heading": law_title,
            })
    if not parents:
        parents = [{"content": body, "heading": law_title}]

    # 合并相邻都偏小的父（多为相邻短章/附则），减少碎片父；不跨大父合并，避免打断语义章节
    i = 0
    while i < len(parents) - 1:
        a, b = parents[i], parents[i + 1]
        if (
            len(a["content"]) < parent_min
            and len(b["content"]) < parent_min
            and len(a["content"]) + len(b["content"]) + 2 <= parent_max
        ):
            a["content"] = f"{a['content']}\n\n{b['content']}"
            del parents[i + 1]
        else:
            i += 1

    for i, p in enumerate(parents):
        p["parent_index"] = i
    return parents

# 按 第X条 切成整条单元
def _split_article_units(text: str) -> list[dict]:
    pieces = _RE_ARTICLE_ONLY.split(text)
    units: list[dict] = []
    lead = (pieces[0] or "").strip()
    if lead:
        units.append({"heading": None, "text": lead})
    for i in range(1, len(pieces), 2):
        head = pieces[i]
        body = pieces[i + 1] if i + 1 < len(pieces) else ""
        joined = f"{head} {body.strip()}".strip() if body.strip() else head
        units.append({"heading": head, "text": joined})
    return units

# 把若干行按 max_len 打包成片段；单行过长时按句子再切
def _pack_lines(lines: list[str], max_len: int) -> list[str]:
    out: list[str] = []
    buf = ""
    for ln in lines:
        if not ln:
            continue
        if len(ln) <= max_len:
            if buf and len(buf) + len(ln) + 1 <= max_len:
                buf = f"{buf}\n{ln}"
            else:
                if buf:
                    out.append(buf)
                buf = ln
            continue
        # 单行超限：按句子切
        if buf:
            out.append(buf)
            buf = ""
        sents = [s.strip() for s in re.split(r"(?<=[。！？；.!?;])\s*", ln) if s.strip()]
        for sent in sents:
            if len(sent) <= max_len:
                if buf and len(buf) + len(sent) + 1 <= max_len:
                    buf = f"{buf}\n{sent}"
                else:
                    if buf:
                        out.append(buf)
                    buf = sent
            else:
                if buf:
                    out.append(buf)
                    buf = ""
                out.append(sent[:max_len])
    if buf:
        out.append(buf)
    return out

# 把一条切成 ≤ child_max 的片段：先按行，再句子兜底
def _split_unit_into_pieces(text: str, child_max: int) -> list[str]:
    text = text.strip()
    if not text:
        return []
    if len(text) <= child_max:
        return [text]
    lines = [ln.strip() for ln in text.split("\n") if ln.strip()]
    if len(lines) > 1:
        return _pack_lines(lines, child_max)
    sents = [s.strip() for s in re.split(r"(?<=[。！？；.!?;])\s*", text) if s.strip()]
    if len(sents) > 1:
        return _pack_lines(sents, child_max)
    return [text[:child_max]]

#把 ≤ child_max 的片段合并成子模块（尽量靠近 child_max，不跨片段截断）
def _pack_children(pieces: list[str], child_max: int) -> list[str]:
    merged: list[str] = []
    buf = ""
    for pc in pieces:
        if not pc:
            continue
        if not buf:
            buf = pc
        elif len(buf) + len(pc) + 2 <= child_max:
            buf = f"{buf}\n\n{pc}"
        else:
            merged.append(buf)
            buf = pc
    if buf:
        merged.append(buf)
    return merged


def chunk_laws(
    text: str,
    file_id: str,
    *,
    parent_min: int = LAW_PARENT_MIN,
    parent_max: int = LAW_PARENT_MAX,
    child_max: int = LAW_CHILD_MAX,
) -> dict:
    """法律库父子文档切分。

    1. 剥目录前部，按真实章/条结构得到正文；
    2. 按章块把整条打包成父模块（约 parent_min..parent_max 字符，整条不跨父）；
    3. 父内再按条/行切成 ≤ child_max 的子模块（不做子间滑动重叠——命中即整父返回，
       父文本已还原跨子上下文，重叠不必要）。

    返回 {"parents": [...], "children": [...]}：
      parents : [{parent_id, file_id, parent_index, heading, content}]  → PG laws
      children: [{content, metadata:{file_id, chunk_index, total_chunks,
                  heading, parent_index, parent_heading}}]               → Chroma / bm25
    """
    text = text.strip()
    if not text:
        return {"parents": [], "children": []}

    law_title = _law_title(text, file_id)
    body = _law_body(text)
    if not body:
        body = text

    parents = _build_law_parents(body, law_title, parent_min, parent_max)

    children: List[Dict[str, Any]] = []
    chunk_index = 0
    for p in parents:
        units = _split_article_units(p["content"])
        pieces: list[str] = []
        for u in units:
            pieces.extend(_split_unit_into_pieces(u["text"], child_max))
        for child_text in _pack_children(pieces, child_max):
            heading = None
            m = _RE_ARTICLE_ONLY.search(child_text)
            if m:
                heading = m.group(1)
            children.append({
                "content": child_text,
                "metadata": {
                    "file_id": file_id,
                    "chunk_index": chunk_index,
                    "total_chunks": 0,
                    "heading": heading,
                    "parent_index": p["parent_index"],
                    "parent_heading": p["heading"],
                },
            })
            chunk_index += 1

    for i, c in enumerate(children):
        c["metadata"]["total_chunks"] = len(children)
        c["metadata"]["chunk_index"] = i

    return {
        "parents": [
            {
                "parent_id": f"{file_id}_P{p['parent_index']}",
                "file_id": file_id,
                "parent_index": p["parent_index"],
                "heading": p["heading"],
                "content": p["content"],
            }
            for p in parents
        ],
        "children": children,
    }


#一站式，解析+切片
def parse_and_chunk(
    file_source: FileSource,
    *,
    chunk_size: int = DEFAULT_CHUNK_SIZE,
    overlap: int = DEFAULT_OVERLAP,
) -> List[Dict[str, Any]]:
    """Convenience: parse a file, then chunk the result."""
    text = parse_file(file_source)
    return chunk_text(text, file_source["file_id"], chunk_size=chunk_size, overlap=overlap)
